from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Protocol
import base64
import logging


logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    text: str
    parser_used: str


class Parser(Protocol):
    def parse(self, content: str) -> ParsedDocument: ...


class PlainTextParser:
    def parse(self, content: str) -> ParsedDocument:
        return ParsedDocument(text=content, parser_used="plain")


class MarkdownParser:
    def parse(self, content: str) -> ParsedDocument:
        return ParsedDocument(text=content, parser_used="markdown")


class PdfParser:
    """PDF parser using pypdf when content is base64 encoded bytes."""

    def parse(self, content: str) -> ParsedDocument:
        try:
            from pypdf import PdfReader  # type: ignore

            raw = base64.b64decode(content)
            reader = PdfReader(BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if text.strip():
                return ParsedDocument(text=text, parser_used="pdf-pypdf")
        except Exception as exc:
            logger.warning("PDF parsing failed; using plain-text fallback", exc_info=exc)
            return ParsedDocument(text=_decode_base64_to_text(content), parser_used="pdf-fallback-plain")
        return ParsedDocument(text=_decode_base64_to_text(content), parser_used="pdf-fallback-plain")


class DocxParser:
    """DOCX parser using python-docx when content is base64 encoded bytes."""

    def parse(self, content: str) -> ParsedDocument:
        try:
            from docx import Document  # type: ignore

            raw = base64.b64decode(content)
            doc = Document(BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            if text.strip():
                return ParsedDocument(text=text, parser_used="docx-python-docx")
        except Exception as exc:
            logger.warning("DOCX parsing failed; using plain-text fallback", exc_info=exc)
            return ParsedDocument(text=_decode_base64_to_text(content), parser_used="docx-fallback-plain")
        return ParsedDocument(text=_decode_base64_to_text(content), parser_used="docx-fallback-plain")


def _decode_base64_to_text(content: str) -> str:
    try:
        raw = base64.b64decode(content)
        return raw.decode("utf-8", errors="ignore")
    except Exception:
        return content


def parse_document_content(content: str, doc_type: str) -> ParsedDocument:
    doc_type_l = (doc_type or "").lower()
    if doc_type_l in {"md", "markdown"}:
        return MarkdownParser().parse(content)
    if doc_type_l == "pdf":
        return PdfParser().parse(content)
    if doc_type_l == "docx":
        return DocxParser().parse(content)
    return PlainTextParser().parse(content)


def parse_document_bytes(content: bytes, doc_type: str) -> ParsedDocument:
    doc_type_l = (doc_type or "").lower()
    b64 = base64.b64encode(content).decode("utf-8")
    return parse_document_content(b64 if doc_type_l in {"pdf", "docx"} else content.decode("utf-8", errors="ignore"), doc_type_l)
